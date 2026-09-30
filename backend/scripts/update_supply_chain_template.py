from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook


ROOT_DIR = Path(__file__).resolve().parents[2]

if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))


TEMPLATE_PATH = (
    ROOT_DIR
    / "data"
    / "templates"
    / "supply_chain_template.xlsx"
)


def normalize(value: object) -> str:
    if value is None:
        return ""

    return (
        str(value)
        .strip()
        .lower()
        .replace(" ", "_")
    )


def update_supply_allocations_sheet(
    workbook,
) -> bool:

    if "Supply_Allocations" not in workbook.sheetnames:
        raise RuntimeError(
            "Sheet 'Supply_Allocations' not found."
        )

    sheet = workbook[
        "Supply_Allocations"
    ]

    headers: dict[str, int] = {}

    for cell in sheet[1]:
        key = normalize(cell.value)

        if key:
            headers[key] = cell.column

    # ---------------------------------------------------------
    # Add spare capacity if missing.
    # ---------------------------------------------------------

    if "spare_capacity_units" not in headers:

        new_column = (
            sheet.max_column + 1
        )

        cell = sheet.cell(
            row=1,
            column=new_column,
        )

        cell.value = (
            "spare_capacity_units"
        )

        headers[
            "spare_capacity_units"
        ] = new_column

        print(
            "[OK] Added "
            "'spare_capacity_units' "
            "column."
        )

    else:
        print(
            "[OK] "
            "'spare_capacity_units' "
            "already exists."
        )

    # ---------------------------------------------------------
    # Existing rows default to zero.
    #
    # We never invent spare capacity.
    # ---------------------------------------------------------

    spare_column = headers[
        "spare_capacity_units"
    ]

    for row in range(
        2,
        sheet.max_row + 1,
    ):

        cell = sheet.cell(
            row=row,
            column=spare_column,
        )

        if cell.value in (
            None,
            "",
        ):
            cell.value = 0.0

    return True


def main() -> None:

    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(
            f"Template not found: {TEMPLATE_PATH}"
        )

    print(
        f"[INFO] Updating template:\n"
        f"{TEMPLATE_PATH}"
    )

    workbook = load_workbook(
        TEMPLATE_PATH
    )

    update_supply_allocations_sheet(
        workbook
    )

    workbook.save(
        TEMPLATE_PATH
    )

    print(
        "[OK] Supply-chain template updated."
    )


if __name__ == "__main__":
    main()